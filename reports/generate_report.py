"""
Sinh báo cáo đồ án (Word) theo đúng cấu trúc yêu cầu của học phần Nhập môn Big Data, dùng ĐỊNH
DẠNG của file mẫu (reports/template/report_template.docx - tạo bởi build_template.py).

Mọi con số, bảng và hình đều lấy từ kết quả thật của pipeline:
- reports/report_stats.json + reports/figures/report/*.png  (sinh bởi reports/make_figures.py)
- reports/metrics_comparison.csv, metrics_before_hour.csv, metrics_after_hour_untuned.csv
- models/thresholds.json, models/tuning_results.json, models/feature_scaler.json

Chạy:  python reports/make_figures.py  &&  python reports/generate_report.py
Kết quả: reports/final_report.docx (mục lục / danh mục được Word cập nhật tự động nếu máy có Word).
"""

import json
import re
import subprocess
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

ROOT_DIR = Path(__file__).resolve().parent.parent
REPORTS_DIR = ROOT_DIR / "reports"
MODELS_DIR = ROOT_DIR / "models"
FIG_DIR = REPORTS_DIR / "figures" / "report"
TEMPLATE_PATH = REPORTS_DIR / "template" / "report_template.docx"
OUTPUT_PATH = REPORTS_DIR / "final_report.docx"

BULLET_NUM_ID = 39  # danh sách gạch đầu dòng của file mẫu
MARKER_TOF_FIGURES = "{{DANH_MUC_HINH}}"
MARKER_TOF_TABLES = "{{DANH_MUC_BANG}}"
TEXT_WIDTH_CM = 16.0


def vn(x, digits: int = 0) -> str:
    s = f"{float(x):,.{digits}f}"
    return s.replace(",", "_").replace(".", ",").replace("_", ".")


def pct(x, digits: int = 2) -> str:
    return f"{vn(float(x) * 100, digits)}%"


# ============================================================================================
# Bộ dựng tài liệu: dùng style của file mẫu (h1-h4, More, Caption, List Paragraph)
# ============================================================================================
class Report:
    def __init__(self, template: Path):
        self.doc = Document(template)
        self.chapter = 0
        self.fig_no: dict[int, int] = {}
        self.tab_no: dict[int, int] = {}
        self.first_chapter = True
        self.captions: list[tuple[str, str, str]] = []  # (loại, nội dung chú thích, bookmark)

    # ---------- tiêu đề ----------
    def chapter_title(self, text: str):
        self.chapter += 1
        p = self.doc.add_paragraph(text, style="h1")
        if not self.first_chapter:
            p.paragraph_format.page_break_before = True
        self.first_chapter = False

    def h2(self, text):
        self.doc.add_paragraph(text, style="h2")

    def h3(self, text):
        self.doc.add_paragraph(text, style="h3")

    def unnumbered(self, text: str):
        """Tiêu đề không đánh số (KẾT LUẬN, TÀI LIỆU THAM KHẢO) - có trong mục lục."""
        p = self.doc.add_paragraph(text, style="More")
        p.paragraph_format.page_break_before = True

    # ---------- đoạn văn ----------
    def _runs(self, p, text: str):
        # **đậm**, _nghiêng_
        for part in re.split(r"(\*\*[^*]+\*\*|__[^_]+__)", text):
            if part.startswith("**") and part.endswith("**"):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("__") and part.endswith("__"):
                p.add_run(part[2:-2]).italic = True
            elif part:
                p.add_run(part)

    def para(self, text: str, align=None):
        p = self.doc.add_paragraph(style="Normal")
        self._runs(p, text)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = 0
        return p

    def formula(self, text: str):
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = 0
        r = p.add_run(text)
        r.italic = True
        return p

    def bullets(self, items: list[str], level: int = 0):
        for it in items:
            p = self.doc.add_paragraph(style="List Paragraph")
            p._p.get_or_add_pPr().append(parse_xml(
                f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/><w:numId w:val="{BULLET_NUM_ID}"/></w:numPr>'))
            self._runs(p, it)

    # ---------- chú thích ----------
    def _caption(self, label: str, counter: dict, text: str):
        c = self.chapter
        counter[c] = counter.get(c, 0) + 1
        n = counter[c]
        seq_id = f"{label}_{c}."
        p = self.doc.add_paragraph(style="Caption")
        p.add_run(f"{label} {c}.")
        for xml in (
            f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>',
            f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> SEQ {seq_id} \\* ARABIC </w:instrText></w:r>',
            f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>',
            f'<w:r {nsdecls("w")}><w:t>{n}</w:t></w:r>',
            f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>',
        ):
            p._p.append(parse_xml(xml))
        p.add_run(f". {text}")
        bookmark = f"_Cap_{len(self.captions) + 1}"
        bid = 9000 + len(self.captions)
        p._p.insert(1, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bid}" w:name="{bookmark}"/>'))
        p._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bid}"/>'))
        self.captions.append((label, f"{label} {c}.{n}. {text}", bookmark))
        return p, f"{label} {c}.{n}"

    # ---------- hình ----------
    def figure(self, name: str, caption: str, width_cm: float = TEXT_WIDTH_CM) -> str:
        path = FIG_DIR / f"{name}.png"
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = 0
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        _, ref = self._caption("Hình", self.fig_no, caption)
        return ref

    def next_figure_ref(self) -> str:
        return f"Hình {self.chapter}.{self.fig_no.get(self.chapter, 0) + 1}"

    def next_table_ref(self) -> str:
        return f"Bảng {self.chapter}.{self.tab_no.get(self.chapter, 0) + 1}"

    # ---------- bảng ----------
    def table(self, caption: str, headers: list[str], rows: list[list], widths_cm: list[float],
              align: list[str] | None = None, font_pt: float = 11.5) -> str:
        cap, ref = self._caption("Bảng", self.tab_no, caption)
        cap.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        tblPr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            + "".join(f'<w:{b} w:val="single" w:sz="6" w:space="0" w:color="auto"/>'
                      for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
            + "</w:tblBorders>"))
        tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        align = align or ["left"] * len(headers)
        for r_idx, values in enumerate([headers] + rows):
            row = t.rows[r_idx]
            if r_idx == 0:
                row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for c_idx, value in enumerate(values):
                cell = row.cells[c_idx]
                cell.width = Cm(widths_cm[c_idx])
                if r_idx == 0:
                    cell._tc.get_or_add_tcPr().append(parse_xml(
                        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="D9D9D9"/>'))
                cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = 0
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                a = "center" if r_idx == 0 else align[c_idx]
                p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                               "right": WD_ALIGN_PARAGRAPH.RIGHT}[a]
                self._runs(p, str(value))
                for run in p.runs:
                    run.font.size = Pt(font_pt)
                    if r_idx == 0:
                        run.bold = True
        # Bảng ngắn giữ nguyên trên 1 trang (không cắt ngang); bảng dài được sang trang và lặp lại
        # dòng tiêu đề.
        for r_idx, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if len(rows) <= 12 and r_idx < len(t.rows) - 1:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        # khoảng trống sau bảng
        spacer = self.doc.add_paragraph(style="Normal")
        spacer.paragraph_format.space_after = Pt(0)
        return ref

    # ---------- danh mục hình/bảng ----------
    def fill_lists(self):
        """Danh mục hình/bảng: mỗi mục = nội dung chú thích + số trang (trường PAGEREF tới bookmark của
        chú thích, Word tự cập nhật). Không dùng trường TOC \\c vì Word chèn thêm khoảng trắng sau số
        thứ tự ("Hình 2.1 . ...")."""
        def esc(x: str) -> str:
            return x.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        for marker, label in ((MARKER_TOF_FIGURES, "Hình"), (MARKER_TOF_TABLES, "Bảng")):
            p_el = next(el for el in self.doc.element.body.iter(qn("w:p"))
                        if Paragraph(el, self.doc).text.strip() == marker)
            anchor = p_el
            for kind, text, bookmark in self.captions:
                if kind != label:
                    continue
                entry = parse_xml(
                    f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="TableofFigures"/>'
                    '<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9344"/></w:tabs>'
                    '<w:ind w:firstLine="0"/></w:pPr>'
                    f'<w:hyperlink w:anchor="{bookmark}" w:history="1">'
                    f'<w:r><w:t xml:space="preserve">{esc(text)}</w:t></w:r><w:r><w:tab/></w:r>'
                    '<w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
                    f'<w:r><w:instrText xml:space="preserve"> PAGEREF {bookmark} \\h </w:instrText></w:r>'
                    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>?</w:t></w:r>'
                    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink></w:p>')
                anchor.addnext(entry)
                anchor = entry
            p_el.getparent().remove(p_el)

    def save(self, path: Path):
        self.fill_lists()
        self.doc.save(path)


# ============================================================================================
# Nội dung
# ============================================================================================
def load_inputs() -> dict:
    data = json.loads((REPORTS_DIR / "report_stats.json").read_text(encoding="utf-8"))
    tuning_path = MODELS_DIR / "tuning_results.json"
    return {
        "s": data["stats"],
        "t": data["tables"],
        "m": pd.read_csv(REPORTS_DIR / "metrics_comparison.csv").set_index("model"),
        "before": pd.read_csv(REPORTS_DIR / "metrics_before_hour.csv").set_index("model"),
        "after": pd.read_csv(REPORTS_DIR / "metrics_after_hour_untuned.csv").set_index("model"),
        "th": json.loads((MODELS_DIR / "thresholds.json").read_text(encoding="utf-8")),
        "tu": json.loads(tuning_path.read_text(encoding="utf-8")) if tuning_path.exists() else {},
        "sc": json.loads((MODELS_DIR / "feature_scaler.json").read_text(encoding="utf-8")),
    }


def chapter1(R: Report, d: dict):
    s, t = d["s"], d["t"]
    R.chapter_title("GIỚI THIỆU ĐỀ TÀI")
    R.para("Chương này trình bày bối cảnh, bài toán cần giải quyết, mục tiêu, phạm vi của đề tài và "
           "giới thiệu chi tiết bộ dữ liệu được sử dụng, bao gồm nguồn gốc, quy mô, cấu trúc và các "
           "đặc trưng liên quan đến tính chất dữ liệu lớn.")

    R.h2("Bối cảnh và lý do chọn đề tài")
    R.para("Thanh toán bằng thẻ tín dụng và thanh toán trực tuyến ngày càng phổ biến, kéo theo số lượng "
           "giao dịch phát sinh mỗi ngày tại các ngân hàng và tổ chức thanh toán lên tới hàng triệu, thậm "
           "chí hàng trăm triệu giao dịch. Cùng với đó, gian lận thẻ (sử dụng thẻ bị đánh cắp, lộ thông tin "
           "thẻ khi mua hàng trực tuyến, thử thẻ bằng các giao dịch nhỏ…) gây thiệt hại tài chính trực tiếp "
           "cho khách hàng và ngân hàng, đồng thời làm suy giảm niềm tin vào hệ thống thanh toán.")
    R.para("Việc kiểm tra thủ công hoặc dựa trên các luật cố định không theo kịp khối lượng giao dịch lớn và "
           "các thủ đoạn gian lận liên tục thay đổi. Bài toán đặt ra yêu cầu một hệ thống có khả năng **lưu "
           "trữ và xử lý dữ liệu giao dịch quy mô lớn một cách phân tán**, đồng thời **học từ dữ liệu lịch sử** "
           "để tự động phát hiện các giao dịch bất thường. Đây là lý do nhóm chọn đề tài **“Phát hiện giao dịch "
           "tài chính bất thường”**, kết hợp công nghệ lưu trữ phân tán Hadoop HDFS, nền tảng xử lý Apache "
           "Spark và các mô hình học máy.")

    R.h2("Phát biểu bài toán")
    R.para("Bài toán được phát biểu dưới dạng **phân lớp nhị phân trên dữ liệu mất cân bằng nặng**: với mỗi "
           "giao dịch thẻ, hệ thống cần dự đoán giao dịch đó là **bình thường (0)** hay **gian lận (1)**.")
    R.bullets([
        "**Đầu vào:** một giao dịch gồm thời điểm giao dịch, số tiền và 28 thuộc tính đã được ngân hàng "
        "ẩn danh hóa bằng phép biến đổi PCA.",
        "**Đầu ra:** xác suất giao dịch là gian lận, nhãn dự đoán theo một ngưỡng quyết định, kèm giải thích "
        "các yếu tố khiến mô hình nghi ngờ.",
        f"**Khó khăn chính:** giao dịch gian lận chỉ chiếm khoảng {pct(s['fraud_rate'])} tổng số giao dịch, "
        "nên một mô hình “đoán tất cả là bình thường” vẫn đạt độ chính xác (accuracy) trên 99,8% nhưng "
        "hoàn toàn vô dụng.",
    ])

    R.h2("Mục tiêu của đề tài")
    R.bullets([
        "Xây dựng quy trình xử lý dữ liệu hoàn chỉnh trên nền tảng Big Data: lưu trữ dữ liệu trên HDFS, "
        "dùng Spark DataFrame để khảo sát, làm sạch, tiền xử lý và tạo đặc trưng.",
        "Xử lý vấn đề mất cân bằng dữ liệu bằng trọng số lớp (class weight) và SMOTE.",
        "Huấn luyện, tinh chỉnh và so sánh ba mô hình: Logistic Regression, Random Forest (Spark MLlib) và "
        "Isolation Forest (scikit-learn).",
        "Đánh giá bằng các độ đo phù hợp với dữ liệu mất cân bằng: Precision, Recall, F1-score, PR-AUC, "
        "ROC-AUC; phân tích ưu, nhược điểm của từng mô hình.",
        "Trực quan hóa dữ liệu, kết quả mô hình và các giao dịch nghi vấn; xây dựng ứng dụng minh họa.",
    ])

    R.h2("Phạm vi và giới hạn")
    R.bullets([
        "Dữ liệu: bộ dữ liệu công khai Credit Card Fraud Detection gồm các giao dịch trong 2 ngày của chủ "
        "thẻ châu Âu (tháng 9/2013).",
        "Xử lý theo lô (batch). Chế độ “giám sát trực tiếp” của ứng dụng là mô phỏng phát lại các giao dịch "
        "của tập kiểm tra theo đúng thứ tự thời gian.",
        "Hệ thống chạy trên một máy tính cá nhân: HDFS dạng pseudo-cluster (1 NameNode, 1 DataNode) bằng "
        "Docker; Spark chạy chế độ local với nhiều luồng song song.",
        "Do 28 thuộc tính V1–V28 đã được ẩn danh, đề tài không diễn giải ý nghĩa nghiệp vụ cụ thể của từng "
        "thuộc tính này.",
    ])

    R.h2("Bộ dữ liệu sử dụng")
    R.h3("Nguồn và quy mô dữ liệu")
    R.para("Đề tài sử dụng bộ dữ liệu **Credit Card Fraud Detection** do nhóm nghiên cứu Machine Learning "
           "Group – Université Libre de Bruxelles (ULB) phối hợp với Worldline công bố trên Kaggle [1], [2]. "
           f"Thông tin tổng quan được trình bày trong {R.next_table_ref()}.")
    R.table("Thông tin tổng quan về bộ dữ liệu", ["Thuộc tính", "Giá trị"], [
        ["Nguồn", "Kaggle – mlg-ulb/creditcardfraud (ULB & Worldline)"],
        ["Định dạng", "CSV (Comma-Separated Values), có dòng tiêu đề"],
        ["Dung lượng", f"{vn(s['file_mb'], 1)} MB"],
        ["Số giao dịch (dòng)", vn(s["raw_rows"])],
        ["Số thuộc tính (cột)", f"{s['raw_cols']} (30 thuộc tính đầu vào + 1 nhãn)"],
        ["Khoảng thời gian", f"2 ngày ({vn(s['time_max_s'])} giây ≈ {vn(s['time_max_s'] / 3600, 1)} giờ)"],
        ["Số giao dịch gian lận", f"{vn(s['raw_fraud'])} ({pct(s['raw_fraud'] / s['raw_rows'], 3)})"],
        ["Giá trị thiếu", vn(s["missing_total"])],
        ["Dòng trùng lặp", vn(s["duplicates"])],
    ], [5.0, 11.0])

    R.h3("Các đặc trưng của bộ dữ liệu")
    R.para("Bộ dữ liệu gồm 31 cột. Vì lý do bảo mật, ngân hàng không công bố các thông tin gốc của giao "
           "dịch (loại thẻ, nơi giao dịch, ngành hàng…) mà chỉ cung cấp 28 thành phần chính V1–V28 thu được "
           "sau phép biến đổi PCA (Principal Component Analysis). Chỉ có hai thuộc tính **Time** và "
           f"**Amount** là giữ nguyên giá trị gốc. Ý nghĩa của từng nhóm đặc trưng được mô tả trong "
           f"{R.next_table_ref()}.")
    R.table("Mô tả các đặc trưng của bộ dữ liệu", ["Đặc trưng", "Kiểu dữ liệu", "Ý nghĩa"], [
        ["Time", "Số thực (giây)", "Số giây tính từ giao dịch đầu tiên trong bộ dữ liệu đến giao dịch hiện tại "
                                   "(0 – 172.792)."],
        ["V1, V2, …, V28", "Số thực", "28 thành phần chính PCA đã ẩn danh, đại diện cho các thông tin gốc của "
                                     "giao dịch; đã được chuẩn hóa quanh giá trị trung bình 0."],
        ["Amount", "Số thực (€)", "Số tiền của giao dịch."],
        ["Class", "Số nguyên (0/1)", "Nhãn: 1 = giao dịch gian lận, 0 = giao dịch bình thường."],
    ], [3.3, 3.2, 9.5])
    R.para(f"{R.next_table_ref()} trình bày thống kê mô tả của toàn bộ 31 đặc trưng do Spark tính trên "
           "toàn bộ dữ liệu (hàm describe). Có thể thấy các thành phần V1–V28 có trung bình xấp xỉ 0 – đặc "
           "điểm của dữ liệu sau PCA – trong khi Time và Amount có thang đo rất khác (Amount lớn nhất "
           f"{vn(t['describe_all']['Amount']['max'], 2)} €), vì vậy cần chuẩn hóa trước khi huấn luyện.")
    rows = []
    for col, v in t["describe_all"].items():
        rows.append([col, vn(v["mean"], 3), vn(v["stddev"], 3), vn(v["min"], 3), vn(v["max"], 3)])
    R.table("Thống kê mô tả các đặc trưng (Spark describe)",
            ["Đặc trưng", "Trung bình", "Độ lệch chuẩn", "Nhỏ nhất", "Lớn nhất"], rows,
            [3.0, 3.2, 3.2, 3.3, 3.3], ["left", "right", "right", "right", "right"], font_pt=10.5)

    R.h3("Đặc điểm dữ liệu lớn của bài toán")
    R.para("Xét theo mô hình 5V của dữ liệu lớn, bài toán phát hiện gian lận thẻ có các đặc điểm sau "
           f"({R.next_table_ref()}). Bộ dữ liệu thực nghiệm có dung lượng vừa phải để chạy trên máy cá nhân, "
           "nhưng giải pháp được thiết kế theo kiến trúc phân tán để có thể mở rộng khi áp dụng cho dữ liệu "
           "thực tế của ngân hàng.")
    R.table("Đặc điểm 5V của bài toán phát hiện gian lận", ["Đặc điểm", "Thể hiện trong bài toán"], [
        ["Volume (khối lượng)", f"{vn(s['raw_rows'])} giao dịch chỉ trong 2 ngày; trong thực tế một ngân hàng "
                                "phát sinh hàng triệu giao dịch mỗi ngày, dữ liệu tích lũy nhiều năm."],
        ["Velocity (tốc độ)", "Giao dịch phát sinh liên tục, cần đánh giá rủi ro gần như tức thời trước khi duyệt."],
        ["Variety (đa dạng)", "Dữ liệu giao dịch có cấu trúc kết hợp thông tin thẻ, thiết bị, vị trí, hành vi "
                              "(trong bộ dữ liệu được nén thành 28 thành phần PCA)."],
        ["Veracity (độ tin cậy)", f"Có {vn(s['duplicates'])} dòng trùng lặp, nhiều giá trị ngoại lệ về số tiền; "
                                  "nhãn gian lận cực kỳ mất cân bằng."],
        ["Value (giá trị)", f"Phát hiện sớm gian lận giúp giảm thiệt hại – riêng 2 ngày dữ liệu đã có "
                            f"{vn(s['fraud_amount_total'], 0)} € bị gian lận."],
    ], [4.0, 12.0])

    R.h2("Kết quả đầu ra và giá trị dự kiến")
    R.bullets([
        "Quy trình xử lý dữ liệu tự động trên HDFS và Spark, chạy lại được bằng một lệnh.",
        "Bộ mô hình đã huấn luyện và tinh chỉnh, cùng bảng so sánh định lượng giữa các mô hình.",
        "Hệ thống biểu đồ trực quan hóa dữ liệu và kết quả, giúp rút ra các thông tin có giá trị như khung "
        "giờ, khoảng số tiền có rủi ro gian lận cao.",
        "Ứng dụng web minh họa: giám sát dòng giao dịch, cảnh báo giao dịch nghi vấn và giải thích lý do.",
    ])

    R.h2("Bố cục báo cáo")
    R.para("Ngoài phần mở đầu, báo cáo gồm các phần: **Chương 1** giới thiệu đề tài và bộ dữ liệu; "
           "**Chương 2** trình bày cơ sở lý thuyết về dữ liệu lớn, Hadoop HDFS, Apache Spark, dữ liệu mất cân "
           "bằng, các thuật toán và độ đo đánh giá; **Chương 3** trình bày quá trình xây dựng ứng dụng: kiến "
           "trúc, lưu trữ, xử lý, khảo sát và trực quan hóa dữ liệu, huấn luyện, đánh giá mô hình; cuối cùng "
           "là phần **Kết luận và hướng phát triển** và tài liệu tham khảo.")


def chapter2(R: Report, d: dict):
    s = d["s"]
    R.chapter_title("CƠ SỞ LÝ THUYẾT")
    R.para("Chương này trình bày các kiến thức nền tảng được sử dụng trong đề tài: khái niệm dữ liệu lớn, "
           "hệ thống lưu trữ phân tán HDFS, nền tảng xử lý Apache Spark, bài toán dữ liệu mất cân bằng, các "
           "thuật toán học máy, độ đo đánh giá và kỹ thuật giải thích mô hình.")

    R.h2("Tổng quan về dữ liệu lớn")
    R.para("Dữ liệu lớn (Big Data) là tập dữ liệu có khối lượng lớn, tốc độ phát sinh nhanh và đa dạng về "
           "định dạng đến mức các công cụ xử lý truyền thống trên một máy đơn lẻ không thể lưu trữ và phân "
           "tích hiệu quả. Dữ liệu lớn thường được mô tả bằng 5 đặc trưng (5V): **Volume** (khối lượng), "
           "**Velocity** (tốc độ), **Variety** (đa dạng), **Veracity** (độ tin cậy) và **Value** (giá trị).")
    R.para("Để xử lý dữ liệu lớn, các hệ thống hiện đại áp dụng hai nguyên tắc: **lưu trữ phân tán** – chia "
           "dữ liệu thành nhiều khối và lưu trên nhiều máy – và **tính toán song song** – đưa phép tính đến nơi "
           "dữ liệu được lưu, chia công việc cho nhiều tiến trình xử lý đồng thời. Hệ sinh thái Hadoop và "
           "Apache Spark là hai nền tảng tiêu biểu hiện thực hai nguyên tắc này.")

    R.h2("Hệ thống tệp phân tán Hadoop HDFS")
    R.h3("Kiến trúc HDFS")
    R.para("HDFS (Hadoop Distributed File System) là hệ thống tệp phân tán của Hadoop, được thiết kế để lưu "
           "trữ các tệp rất lớn trên cụm máy tính thông thường với khả năng chịu lỗi cao [3]. HDFS có kiến trúc "
           f"master/slave như {R.next_figure_ref()}:")
    R.bullets([
        "**NameNode** (master): quản lý không gian tên của hệ thống tệp (cây thư mục, quyền truy cập) và "
        "metadata cho biết mỗi tệp gồm những block nào, các block nằm trên DataNode nào.",
        "**DataNode** (slave): lưu trữ dữ liệu thực tế dưới dạng các block, định kỳ gửi tín hiệu "
        "heartbeat và báo cáo danh sách block cho NameNode.",
        "**Block:** mỗi tệp được chia thành các block kích thước cố định (mặc định 128 MB). Mỗi block được "
        "nhân bản (replication, mặc định 3 bản) trên nhiều DataNode để đảm bảo không mất dữ liệu khi một "
        "máy gặp sự cố.",
    ])
    R.figure("hdfs_nguyen_ly", "Kiến trúc và cơ chế lưu trữ block của HDFS", 15)
    R.h3("Cơ chế đọc và ghi dữ liệu")
    R.para("Khi đọc tệp, client hỏi NameNode vị trí các block, sau đó đọc trực tiếp từ các DataNode chứa "
           "block đó; NameNode không nằm trên đường truyền dữ liệu nên không trở thành nút thắt cổ chai. Khi "
           "ghi, dữ liệu được ghi theo đường ống (pipeline) qua các DataNode giữ bản sao. HDFS tối ưu cho mô "
           "hình “ghi một lần, đọc nhiều lần” – rất phù hợp với dữ liệu giao dịch lịch sử dùng để phân tích "
           f"và huấn luyện mô hình. Với đề tài này, tệp {vn(s['file_mb'])} MB được HDFS chia thành "
           f"{s['hdfs_blocks']} block.")

    R.h2("Nền tảng xử lý dữ liệu Apache Spark")
    R.h3("Tổng quan")
    R.para("Apache Spark là nền tảng xử lý dữ liệu phân tán mã nguồn mở, ra đời để khắc phục hạn chế của "
           "Hadoop MapReduce [4], [5]. Thay vì ghi kết quả trung gian xuống đĩa sau mỗi bước như MapReduce, "
           "Spark giữ dữ liệu trong bộ nhớ (in-memory) giữa các bước tính toán, nhờ đó nhanh hơn nhiều với các "
           "thuật toán lặp như huấn luyện mô hình học máy. Spark cung cấp các thư viện thống nhất: Spark SQL/"
           "DataFrame, Spark Streaming, MLlib (học máy) và GraphX (đồ thị); có thể đọc dữ liệu trực tiếp từ "
           "HDFS.")
    R.h3("RDD, DataFrame và cơ chế đánh giá lười")
    R.bullets([
        "**RDD (Resilient Distributed Dataset):** tập dữ liệu phân tán, bất biến, được chia thành nhiều "
        "partition xử lý song song; có khả năng tự khôi phục nhờ lưu lại chuỗi phép biến đổi (lineage) [4].",
        "**DataFrame:** dữ liệu dạng bảng có lược đồ (schema) gồm các cột có tên, xây dựng trên RDD. "
        "DataFrame được bộ tối ưu Catalyst tối ưu kế hoạch thực thi, cho hiệu năng cao và cú pháp gần với SQL. "
        "Đề tài sử dụng DataFrame cho toàn bộ bước xử lý dữ liệu.",
        "**Đánh giá lười (lazy evaluation):** các phép biến đổi (transformation) như select, filter, "
        "withColumn chỉ ghi nhận kế hoạch; tính toán thật chỉ diễn ra khi gặp hành động (action) như count, "
        "collect, write.",
    ])
    R.h3("Mô hình thực thi song song")
    R.para(f"Mô hình thực thi của Spark được minh họa trong {R.next_figure_ref()}. Chương trình driver tạo "
           "SparkSession; mỗi action sinh ra một **job**; job được chia thành các **stage** tại ranh giới "
           "shuffle (khi dữ liệu cần trao đổi giữa các partition, ví dụ groupBy, dropDuplicates); mỗi stage gồm "
           "nhiều **task**, mỗi task xử lý một **partition** dữ liệu. Các task được phân cho các executor chạy "
           f"song song. Trong đề tài, Spark chạy chế độ local[{s['spark_cores']}], tức {s['spark_cores']} luồng "
           "xử lý song song trên một máy; cùng một mã nguồn có thể chạy trên cụm nhiều máy chỉ bằng cách đổi "
           "địa chỉ master.")
    R.figure("spark_mo_hinh_thuc_thi", "Mô hình thực thi job – stage – task của Spark", 15)
    R.h3("Thư viện học máy Spark MLlib")
    R.para("Spark MLlib cung cấp các thuật toán học máy chạy phân tán và các công cụ xây dựng quy trình: "
           "**VectorAssembler** (gộp các cột thành vector đặc trưng), **StandardScaler** (chuẩn hóa), các bộ "
           "phân lớp **LogisticRegression**, **RandomForestClassifier** hỗ trợ cột trọng số mẫu (weightCol), "
           "**BinaryClassificationEvaluator** (tính ROC-AUC, PR-AUC) và **CrossValidator** + "
           "**ParamGridBuilder** cho tìm kiếm siêu tham số.")
    R.h3("Lý do lựa chọn HDFS và Spark")
    R.para(f"{R.next_table_ref()} so sánh giải pháp được chọn với cách xử lý truyền thống trên một máy.")
    R.table("So sánh giải pháp Big Data với cách xử lý truyền thống",
            ["Tiêu chí", "Truyền thống (tệp cục bộ + pandas)", "HDFS + Spark (đề tài)"], [
                ["Lưu trữ", "Một tệp trên một ổ đĩa, giới hạn bởi dung lượng máy, mất dữ liệu khi hỏng ổ",
                 "Chia block, lưu phân tán và nhân bản trên nhiều DataNode, chịu lỗi"],
                ["Xử lý", "Một luồng, toàn bộ dữ liệu phải nạp vào RAM",
                 "Song song trên nhiều partition/luồng/máy, xử lý được dữ liệu lớn hơn RAM"],
                ["Mở rộng", "Nâng cấp máy (scale-up), có giới hạn",
                 "Thêm máy vào cụm (scale-out), mã nguồn không đổi"],
                ["Học máy", "scikit-learn trên một máy", "MLlib huấn luyện phân tán, tích hợp cùng pipeline"],
                ["Phù hợp", "Dữ liệu nhỏ, phân tích nhanh", "Dữ liệu giao dịch tích lũy liên tục, khối lượng lớn"],
            ], [3.0, 6.5, 6.5])

    R.h2("Bài toán dữ liệu mất cân bằng")
    R.para("Dữ liệu mất cân bằng (imbalanced data) là dữ liệu trong đó một lớp (lớp thiểu số – gian lận) "
           "có số mẫu ít hơn rất nhiều so với lớp còn lại. Khi đó các mô hình có xu hướng thiên về lớp đa "
           "số. Đề tài sử dụng hai kỹ thuật:")
    R.bullets([
        "**Trọng số lớp (class weight):** gán cho mỗi mẫu một trọng số tỉ lệ nghịch với tần suất lớp của nó, "
        "để lỗi phân loại sai một mẫu gian lận bị “phạt” nặng hơn nhiều so với một mẫu bình thường. Cách này "
        "không làm thay đổi số lượng dữ liệu.",
        "**SMOTE (Synthetic Minority Over-sampling Technique) [6]:** sinh thêm mẫu thiểu số nhân tạo bằng "
        "cách nội suy giữa một mẫu thiểu số và một trong k láng giềng gần nhất cùng lớp của nó.",
    ])
    R.formula("w(c) = N / (K · N_c)        x_mới = x_i + λ · (x_láng_giềng − x_i),  λ ∈ [0, 1]")
    R.para("trong đó N là tổng số mẫu, K là số lớp, N_c là số mẫu của lớp c.", align=None)

    R.h2("Các thuật toán sử dụng")
    R.h3("Logistic Regression")
    R.para("Logistic Regression là mô hình tuyến tính ước lượng xác suất một mẫu thuộc lớp dương bằng hàm "
           "sigmoid của tổ hợp tuyến tính các đặc trưng. Mô hình được huấn luyện bằng cách cực tiểu hóa hàm "
           "mất mát log-loss (có trọng số mẫu) cộng thành phần điều chuẩn Elastic Net – tổ hợp của L1 (giúp "
           "chọn đặc trưng) và L2 (giảm quá khớp), điều khiển bởi hai tham số regParam (λ) và "
           "elasticNetParam (α). Ưu điểm: đơn giản, nhanh, dễ giải thích; nhược điểm: chỉ học được ranh giới "
           "quyết định tuyến tính.")
    R.formula("P(y = 1 | x) = σ(wᵀx + b) = 1 / (1 + e^−(wᵀx + b))")
    R.h3("Random Forest")
    R.para("Random Forest [7] là phương pháp học tập hợp (ensemble) gồm nhiều cây quyết định. Mỗi cây được "
           "huấn luyện trên một mẫu bootstrap của dữ liệu và tại mỗi nút chỉ xét một tập con ngẫu nhiên các "
           "đặc trưng để chọn điểm tách (theo độ đo Gini). Xác suất dự đoán là trung bình xác suất của các "
           "cây. Nhờ tính ngẫu nhiên và lấy trung bình, Random Forest giảm phương sai, chống quá khớp, học "
           "được quan hệ phi tuyến và tương tác giữa các đặc trưng. Hai siêu tham số chính là số cây "
           "(numTrees) và độ sâu tối đa (maxDepth).")
    R.h3("Isolation Forest")
    R.para("Isolation Forest [8] là thuật toán phát hiện bất thường **không giám sát** – không dùng nhãn khi "
           "huấn luyện. Ý tưởng: điểm bất thường “ít và khác biệt” nên dễ bị cô lập. Thuật toán xây dựng nhiều "
           "cây cô lập bằng cách chọn ngẫu nhiên đặc trưng và ngưỡng tách; độ dài đường đi trung bình h(x) từ "
           "gốc đến khi cô lập được điểm x càng ngắn thì điểm đó càng bất thường. Điểm bất thường được tính:")
    R.formula("s(x, n) = 2^( −E[h(x)] / c(n) )")
    R.para("trong đó c(n) là độ dài đường đi trung bình chuẩn hóa với n mẫu; s gần 1 nghĩa là rất bất "
           "thường. Tham số contamination xác định tỉ lệ điểm được coi là bất thường.")

    R.h2("Tạo đặc trưng thời gian dạng chu kỳ")
    R.para("Các đại lượng có tính chu kỳ như giờ trong ngày không nên biểu diễn trực tiếp bằng số 0–23, vì "
           "khi đó mô hình coi 23h và 0h cách nhau rất xa trong khi thực tế chỉ cách 1 giờ. Cách phổ biến là "
           "ánh xạ giờ lên đường tròn đơn vị bằng hai đặc trưng sin và cos (tương tự vị trí kim đồng hồ):")
    R.formula("Hour_sin = sin(2π · h / 24),        Hour_cos = cos(2π · h / 24)")
    R.para("Với cách biểu diễn này, các thời điểm gần nhau trong ngày có tọa độ gần nhau, và một mô hình "
           "tuyến tính có thể học được quy luật dạng “ban đêm rủi ro cao” vốn không đơn điệu theo giờ.")

    R.h2("Các độ đo đánh giá mô hình")
    R.para(f"Kết quả phân lớp nhị phân được tổng hợp trong ma trận nhầm lẫn ({R.next_table_ref()}), từ đó "
           "tính các độ đo.")
    R.table("Ma trận nhầm lẫn của bài toán phân lớp nhị phân",
            ["", "Dự đoán: gian lận", "Dự đoán: bình thường"], [
                ["Thực tế: gian lận", "TP (phát hiện đúng)", "FN (bỏ sót)"],
                ["Thực tế: bình thường", "FP (báo nhầm)", "TN (duyệt đúng)"],
            ], [4.6, 5.7, 5.7], ["left", "center", "center"])
    R.bullets([
        "**Precision = TP / (TP + FP):** trong các giao dịch bị cảnh báo, bao nhiêu phần trăm thật sự là "
        "gian lận – phản ánh chi phí báo nhầm (làm phiền khách hàng).",
        "**Recall = TP / (TP + FN):** trong các giao dịch gian lận, mô hình phát hiện được bao nhiêu phần "
        "trăm – phản ánh thiệt hại do bỏ sót.",
        "**F1-score = 2 · Precision · Recall / (Precision + Recall):** trung bình điều hòa, cân bằng hai độ "
        "đo trên.",
        "**ROC-AUC:** diện tích dưới đường cong ROC (TPR theo FPR) khi thay đổi ngưỡng.",
        "**PR-AUC (Average Precision):** diện tích dưới đường cong Precision–Recall.",
    ])
    R.para("Với dữ liệu mất cân bằng nặng, **Accuracy** gần như vô nghĩa và **ROC-AUC** thường cho giá trị "
           "lạc quan vì FPR được tính trên số lượng rất lớn mẫu âm. **PR-AUC** chỉ tập trung vào lớp gian "
           "lận nên phản ánh đúng hơn chất lượng mô hình [9], [10]. Vì vậy đề tài dùng PR-AUC làm độ đo chính "
           "để tinh chỉnh và chọn mô hình.")

    R.h2("Tinh chỉnh siêu tham số và ngưỡng phân loại")
    R.bullets([
        "**Grid Search + k-fold Cross-Validation:** thử mọi tổ hợp tham số trong lưới; với mỗi tổ hợp, chia "
        "tập huấn luyện thành k phần, lần lượt huấn luyện trên k−1 phần và đánh giá trên phần còn lại, lấy "
        "trung bình độ đo; chọn tổ hợp tốt nhất.",
        "**Tối ưu ngưỡng phân loại:** mô hình xuất ra xác suất; ngưỡng mặc định 0,5 thường không tối ưu với "
        "dữ liệu mất cân bằng. Ngưỡng được chọn để F1-score lớn nhất trên một **tập validation tách riêng** "
        "(không dùng tập test, tránh rò rỉ dữ liệu).",
        "**Chia tập theo thời gian (time-based split):** giao dịch cũ dùng để huấn luyện, giao dịch mới dùng "
        "để kiểm tra – mô phỏng đúng thực tế dự đoán giao dịch tương lai.",
    ])

    R.h2("Giải thích mô hình với SHAP")
    R.para("Để người dùng tin tưởng một cảnh báo gian lận, hệ thống cần giải thích được vì sao mô hình đưa ra "
           "kết luận đó (Explainable AI – xAI). SHAP (SHapley Additive exPlanations) [11] dựa trên giá trị "
           "Shapley của lý thuyết trò chơi: mỗi đặc trưng nhận một giá trị đóng góp sao cho **giá trị cơ sở "
           "(trung bình) cộng tổng các đóng góp đúng bằng đầu ra của mô hình** cho giao dịch đó. Với các mô "
           "hình dạng cây như Random Forest, thuật toán TreeSHAP [12] tính chính xác các giá trị này trong thời "
           "gian đa thức. Tính chất cộng dồn cho phép gộp đóng góp của nhiều đặc trưng thành một nhóm lý do dễ "
           "hiểu (ví dụ gộp V1–V28 thành “hành vi thẻ”).")


def chapter3(R: Report, d: dict):
    s, t, m, th, tu, sc = d["s"], d["t"], d["m"], d["th"], d["tu"], d["sc"]
    rf, lr, iso = m.loc["Random Forest"], m.loc["Logistic Regression"], m.loc["Isolation Forest"]
    R.chapter_title("XÂY DỰNG ỨNG DỤNG, TRỰC QUAN HÓA DỮ LIỆU")
    R.para("Chương này trình bày quá trình hiện thực giải pháp: kiến trúc hệ thống, môi trường thực nghiệm, "
           "lưu trữ dữ liệu trên HDFS, khảo sát và trực quan hóa dữ liệu bằng Spark, làm sạch và tiền xử lý, "
           "huấn luyện, tinh chỉnh, đánh giá mô hình và ứng dụng minh họa.")

    # ---- 3.1 Kiến trúc
    R.h2("Kiến trúc tổng thể và luồng dữ liệu")
    R.para(f"Kiến trúc hệ thống được thể hiện trong {R.next_figure_ref()}. Dữ liệu thô được lưu trên HDFS; "
           "Spark đọc dữ liệu từ HDFS để khảo sát, làm sạch, tạo đặc trưng và chuẩn hóa; kết quả trung gian "
           "được ghi lại dưới dạng Parquet (định dạng cột, nén tốt, đọc nhanh). Tiếp theo là các bước xử lý "
           "mất cân bằng, huấn luyện, tinh chỉnh và đánh giá mô hình; cuối cùng là trực quan hóa và ứng dụng web.")
    R.figure("kien_truc_he_thong", "Kiến trúc tổng thể và luồng dữ liệu của hệ thống")
    R.para(f"Mã nguồn được tổ chức thành các mô-đun độc lập theo từng giai đoạn ({R.next_table_ref()}); "
           "tệp run_pipeline.py điều phối chạy toàn bộ quy trình.")
    R.table("Các mô-đun chính của hệ thống", ["Tệp mã nguồn", "Chức năng"], [
        ["src/config.py", "Cấu hình đường dẫn (local/HDFS), danh sách đặc trưng, tham số chung"],
        ["src/spark_etl.py", "Đọc dữ liệu từ HDFS, EDA, làm sạch, phân tích ngoại lệ, tạo đặc trưng giờ, chuẩn "
                             "hóa, chia tập theo thời gian, ghi Parquet"],
        ["src/imbalance.py", "Tính trọng số lớp, sinh dữ liệu SMOTE"],
        ["src/train_models.py", "Huấn luyện Logistic Regression, Random Forest (MLlib), Isolation Forest"],
        ["src/tune_models.py", "Grid Search + Cross-Validation, so sánh class weight/SMOTE, tối ưu ngưỡng"],
        ["src/evaluate.py", "Đánh giá trên tập test, xuất bảng so sánh"],
        ["src/visualize.py, reports/make_figures.py", "Trực quan hóa dữ liệu và kết quả bằng matplotlib"],
        ["src/hdfs_utils.py, scripts/hdfs_bootstrap.py", "Thao tác HDFS qua WebHDFS, đưa dữ liệu lên HDFS"],
        ["run_pipeline.py", "Điều phối toàn bộ quy trình"],
    ], [6.0, 10.0])

    # ---- 3.2 Môi trường
    R.h2("Môi trường thực nghiệm")
    R.table("Môi trường phần cứng và phần mềm", ["Thành phần", "Cấu hình / phiên bản"], [
        ["CPU", "Intel Core i5-12450HX (8 nhân, 12 luồng)"],
        ["Bộ nhớ", "28 GB RAM"],
        ["Hệ điều hành", "Windows 11"],
        ["Ngôn ngữ", "Python 3.11.9"],
        ["Xử lý phân tán", f"Apache Spark (PySpark) 4.2.0, chế độ local[{s['spark_cores']}], Java 17"],
        ["Lưu trữ phân tán", "Hadoop HDFS 3.2.1 (Docker image bde2020), 1 NameNode + 1 DataNode"],
        ["Học máy", "Spark MLlib, scikit-learn 1.9.1, imbalanced-learn 0.14.2"],
        ["Xử lý dữ liệu, trực quan", "pandas 2.3.3, NumPy 2.4.6, matplotlib 3.11.2"],
        ["Ứng dụng web", "FastAPI (backend), Next.js (frontend)"],
    ], [4.5, 11.5])

    # ---- 3.3 HDFS
    R.h2("Lưu trữ dữ liệu trên HDFS")
    R.para("HDFS được triển khai dạng pseudo-cluster bằng Docker Compose (tệp docker-compose.hdfs.yml) gồm "
           "một container NameNode (cổng giao diện web 9870, cổng RPC 9000) và một container DataNode. Dữ liệu "
           "thô được đưa lên HDFS qua WebHDFS REST API bằng script scripts/hdfs_bootstrap.py, theo cấu trúc thư "
           "mục:")
    R.bullets([
        "/data/raw/creditcard.csv – dữ liệu thô.",
        "/data/processed/cleaned.parquet, train.parquet, test.parquet, train_weighted.parquet, "
        "train_balanced.parquet – dữ liệu sau xử lý do Spark ghi ra.",
    ])
    R.para("Mã nguồn hỗ trợ hai chế độ thông qua biến môi trường DATA_BACKEND: **hdfs** (Spark đọc/ghi theo "
           "URI hdfs://namenode:9000/data/…) và **local** (thư mục data/ trên máy, dùng khi phát triển). "
           "Cùng một mã xử lý chạy được ở cả hai chế độ mà không cần sửa đổi.")

    # ---- 3.4 EDA
    R.h2("Khảo sát và trực quan hóa dữ liệu")
    R.para("Toàn bộ số liệu khảo sát dưới đây do Spark tổng hợp trên toàn bộ dữ liệu (groupBy, "
           "approxQuantile, Correlation của MLlib); chỉ các kết quả tổng hợp nhỏ được chuyển về để vẽ biểu đồ "
           "bằng matplotlib.")
    R.h3("Phân bố nhãn")
    R.para(f"{R.next_figure_ref()} cho thấy mức độ mất cân bằng: sau khi làm sạch còn {vn(s['clean_rows'])} "
           f"giao dịch, trong đó chỉ {vn(s['clean_fraud'])} giao dịch gian lận ({pct(s['fraud_rate'], 3)}), "
           f"tức cứ khoảng {vn(1 / s['fraud_rate'])} giao dịch mới có một giao dịch gian lận.")
    R.figure("phan_bo_nhan", "Phân bố giao dịch bình thường và gian lận")

    amt = pd.DataFrame(t["amount_buckets"]).set_index("label")
    top_rate = amt["rate"].idxmax()
    top_loss = amt["fraud_amount"].idxmax()
    R.h3("Gian lận theo số tiền giao dịch")
    R.para(f"{R.next_figure_ref()} trình bày tỉ lệ gian lận và tổng số tiền bị gian lận theo từng khoảng số "
           f"tiền. Khoảng **{top_rate} €** có tỉ lệ gian lận cao nhất ({pct(amt.loc[top_rate, 'rate'])}, so với "
           f"mức chung {pct(s['fraud_rate'])}), trong khi khoảng **{top_loss} €** chiếm nhiều tiền gian lận nhất "
           f"({vn(amt.loc[top_loss, 'fraud_amount'], 0)} €). Trung vị số tiền của giao dịch gian lận là "
           f"{vn(s['amount_median_by_class']['1'], 2)} €, thấp hơn giao dịch bình thường "
           f"({vn(s['amount_median_by_class']['0'], 2)} €) – phù hợp với thủ đoạn “thử thẻ” bằng các khoản "
           "nhỏ trước khi thực hiện giao dịch lớn. Như vậy số tiền một mình không đủ để phát hiện gian lận.")
    R.figure("gian_lan_theo_so_tien", "Tỉ lệ gian lận và tổng tiền gian lận theo khoảng số tiền")

    hours = pd.DataFrame(t["hours"]).set_index("Hour")
    peak_h = int(hours["rate"].idxmax())
    day_rate = hours.loc[9:21, "fraud"].sum() / hours.loc[9:21, "n"].sum()
    R.h3("Gian lận theo giờ trong ngày")
    R.para("Cột Time chỉ là số giây kể từ giao dịch đầu tiên. Từ cột này, nhóm suy ra giờ trong ngày với giả "
           "định giao dịch đầu tiên diễn ra lúc 0h; giả định này được củng cố bởi dữ liệu: số giao dịch thấp "
           f"nhất rơi vào khoảng 1h–6h, đúng nhịp sinh hoạt ban đêm ({R.next_figure_ref()}). Kết quả cho thấy "
           f"**lúc {peak_h}h tỉ lệ gian lận lên tới {pct(hours.loc[peak_h, 'rate'])}**, cao gấp khoảng "
           f"{vn(hours.loc[peak_h, 'rate'] / day_rate, 0)} lần mức trung bình ban ngày 9h–21h ({pct(day_rate)}): "
           "ban đêm ít giao dịch hợp lệ nhưng kẻ gian vẫn hoạt động.")
    R.figure("gian_lan_theo_gio", "Số giao dịch và tỉ lệ gian lận theo giờ trong ngày")
    R.para(f"Kết hợp hai yếu tố, {R.next_figure_ref()} cho thấy các tổ hợp khung giờ – khoảng số tiền có "
           "rủi ro cao nhất; đây là kiểu tương tác mà các mô hình dạng cây có thể học được.")
    R.figure("gio_x_so_tien", "Tỉ lệ gian lận theo khung giờ và khoảng số tiền")

    cwc = t["corr_with_class"]
    top4 = s["top_corr_features"]
    R.h3("Quan hệ giữa các đặc trưng")
    R.para(f"{R.next_figure_ref()} liệt kê các đặc trưng có tương quan Pearson mạnh nhất với nhãn: "
           + ", ".join(f"{k} ({vn(cwc[k], 3)})" for k in top4)
           + ". Hệ số âm nghĩa là giá trị đặc trưng càng thấp thì khả năng gian lận càng cao.")
    R.figure("tuong_quan_voi_nhan", "Các đặc trưng tương quan mạnh nhất với nhãn gian lận", 14)
    R.para(f"Phân bố của 4 đặc trưng này ở hai nhóm ({R.next_figure_ref()}) tách biệt rõ rệt: giao dịch "
           "bình thường tập trung quanh 0 trong khi giao dịch gian lận lệch hẳn về một phía.")
    R.figure("phan_bo_dac_trung_manh", "Phân bố của 4 đặc trưng tương quan mạnh nhất theo nhãn")
    R.para(f"{R.next_figure_ref()} so sánh ma trận tương quan giữa các đặc trưng trong từng nhóm. Với giao "
           "dịch bình thường, các thành phần V1–V28 gần như không tương quan với nhau – đúng với tính chất của "
           "PCA. Ngược lại, trong nhóm giao dịch gian lận xuất hiện nhiều cặp đặc trưng tương quan mạnh: các "
           "giao dịch gian lận có một “khuôn mẫu” chung, đây là cơ sở để mô hình học máy phân biệt được hai "
           "nhóm.")
    R.figure("tuong_quan_dac_trung", "Ma trận tương quan giữa các đặc trưng: bình thường và gian lận")

    # ---- 3.5 Làm sạch & tiền xử lý
    R.h2("Làm sạch và tiền xử lý dữ liệu")
    R.h3("Dữ liệu thiếu và trùng lặp")
    R.para(f"Kiểm tra bằng Spark cho thấy dữ liệu không có giá trị thiếu ({vn(s['missing_total'])} ô trống). "
           f"Có {vn(s['duplicates'])} dòng trùng lặp hoàn toàn (trong đó {vn(s['fraud_duplicates_removed'])} "
           "dòng mang nhãn gian lận); các dòng này bị loại bằng dropDuplicates để mô hình không học lặp lại "
           f"cùng một mẫu và tránh rò rỉ giữa tập train và test. Sau làm sạch còn {vn(s['clean_rows'])} giao dịch.")
    o = s["outliers"]
    R.h3("Phân tích ngoại lệ")
    R.para(f"Ngoại lệ của số tiền được xác định theo quy tắc IQR (Tukey): Q1 = {vn(o['q1'], 2)} €, "
           f"Q3 = {vn(o['q3'], 2)} €, IQR = {vn(o['iqr'], 2)} €, ngưỡng trên Q3 + 1,5·IQR = "
           f"{vn(o['upper'], 2)} € ({R.next_figure_ref()}, {R.next_table_ref()}).")
    R.figure("ngoai_le_so_tien", "Phân bố số tiền và ngưỡng ngoại lệ IQR")
    oc = o["by_class"]
    R.table("Thống kê ngoại lệ số tiền theo nhãn", ["Nhóm", "Số giao dịch", "Số ngoại lệ", "Tỉ lệ ngoại lệ"], [
        ["Bình thường", vn(oc["0"]["n"]), vn(oc["0"]["outliers"]), pct(oc["0"]["outliers"] / oc["0"]["n"])],
        ["Gian lận", vn(oc["1"]["n"]), vn(oc["1"]["outliers"]), pct(oc["1"]["outliers"] / oc["1"]["n"])],
    ], [4.0, 4.0, 4.0, 4.0], ["left", "right", "right", "right"])
    R.para(f"Tỉ lệ ngoại lệ ở nhóm gian lận ({pct(oc['1']['outliers'] / oc['1']['n'])}) cao hơn nhóm bình "
           f"thường ({pct(oc['0']['outliers'] / oc['0']['n'])}). Vì vậy nhóm **quyết định giữ lại các ngoại lệ**: "
           "giao dịch gian lận vốn là các điểm bất thường, loại bỏ ngoại lệ sẽ làm mất chính tín hiệu cần học. "
           "Ảnh hưởng của biên độ lớn được xử lý bằng chuẩn hóa (với Logistic Regression), còn Random Forest và "
           "Isolation Forest vốn ít nhạy với ngoại lệ.")
    R.h3("Tạo đặc trưng giờ trong ngày")
    R.para(f"Nhóm tạo thêm các đặc trưng từ cột Time như {R.next_table_ref()}. Đặc trưng Time_scaled (Time "
           "chuẩn hóa) bị **loại khỏi mô hình** vì Time chỉ tăng dần theo thời gian; khi chia tập theo thời gian, "
           "mọi giá trị Time của tập test đều lớn hơn tập train nên mô hình phải dự đoán trên vùng giá trị chưa "
           "từng thấy. Thông tin thời gian được đưa vào mô hình qua cặp Hour_sin, Hour_cos.")
    R.table("Các đặc trưng được tạo thêm", ["Đặc trưng", "Công thức", "Mục đích"], [
        ["Hour", "⌊(Time mod 86400) / 3600⌋", "Giờ trong ngày (0–23), dùng hiển thị, thống kê"],
        ["Hour_sin", "sin(2π·h/24), h = (Time mod 86400)/3600", "Đặc trưng giờ dạng chu kỳ cho mô hình"],
        ["Hour_cos", "cos(2π·h/24)", "Đặc trưng giờ dạng chu kỳ cho mô hình"],
        ["Amount_scaled", "(Amount − μ) / σ", "Số tiền đã chuẩn hóa cho mô hình"],
    ], [3.2, 6.6, 6.2])
    R.h3("Chuẩn hóa dữ liệu")
    R.para("Amount được chuẩn hóa bằng StandardScaler của Spark MLlib về phân phối có trung bình 0, độ lệch "
           f"chuẩn 1, với tham số học được μ = {vn(sc['amount_mean'], 4)} € và σ = {vn(sc['amount_std'], 4)} €. "
           "Tham số được lưu lại (models/feature_scaler.json) để ứng dụng dùng đúng phép biến đổi khi suy luận. "
           "Bộ đặc trưng cuối cùng đưa vào mô hình gồm 31 đặc trưng: Hour_sin, Hour_cos, Amount_scaled và "
           "V1–V28.")
    R.h3("Chia tập huấn luyện và kiểm tra theo thời gian")
    R.para(f"Dữ liệu được chia theo thời gian với tỉ lệ 70/30 ({R.next_figure_ref()}): {vn(s['train_rows'])} "
           f"giao dịch đầu tiên ({vn(s['train_fraud'])} gian lận) để huấn luyện, {vn(s['test_rows'])} giao dịch "
           f"sau mốc giây {vn(s['split_cutoff_s'])} ({vn(s['test_fraud'])} gian lận) để kiểm tra. Cách chia này "
           "mô phỏng đúng tình huống thực tế: mô hình học từ quá khứ để dự đoán giao dịch tương lai.")
    R.figure("chia_tap_theo_thoi_gian", "Chia tập huấn luyện và kiểm tra theo thời gian")

    # ---- 3.6 Mất cân bằng
    R.h2("Xử lý dữ liệu mất cân bằng")
    cw = s["class_weight"]
    R.para(f"Theo công thức trọng số lớp, mỗi giao dịch gian lận có trọng số **{vn(cw['1'], 2)}** và mỗi "
           f"giao dịch bình thường có trọng số **{vn(cw['0'], 4)}** – lỗi bỏ sót một giao dịch gian lận bị phạt "
           f"nặng gấp khoảng {vn(cw['1'] / cw['0'])} lần. Trọng số này được đưa vào tham số weightCol của "
           "LogisticRegression và RandomForestClassifier.")
    R.para(f"Song song, SMOTE được áp dụng với tỉ lệ lớp thiểu số/đa số bằng 0,1 – số mẫu gian lận tăng từ "
           f"{vn(s['train_fraud'])} lên {vn(s['smote']['1'])} ({R.next_figure_ref()}). Không cân bằng hoàn toàn "
           "1:1 vì sinh quá nhiều mẫu nhân tạo dễ gây quá khớp. Hai chiến lược được so sánh ở bước tinh chỉnh.")
    R.figure("smote", "Số mẫu huấn luyện trước và sau SMOTE", 13)

    # ---- 3.7 Huấn luyện
    R.h2("Huấn luyện mô hình")
    R.para(f"Ba mô hình được huấn luyện với cấu hình ban đầu như {R.next_table_ref()}.")
    R.table("Cấu hình huấn luyện ban đầu của các mô hình", ["Mô hình", "Thư viện", "Tham số", "Ghi chú"], [
        ["Logistic Regression", "Spark MLlib", "maxIter = 100, regParam = 0,01", "weightCol = class_weight"],
        ["Random Forest", "Spark MLlib", "numTrees = 100, maxDepth = 8", "weightCol = class_weight"],
        ["Isolation Forest", "scikit-learn", "n_estimators = 200, contamination = 0,01",
         "Không dùng nhãn; huấn luyện trên 50.000 mẫu lấy từ Spark"],
    ], [3.6, 2.6, 4.6, 5.2])
    R.para("Isolation Forest không có phiên bản trong Spark MLlib nên được huấn luyện bằng scikit-learn trên "
           "một mẫu dữ liệu lấy từ Spark DataFrame. Tham số contamination được nới lên 1% (thay vì đúng bằng tỉ "
           "lệ gian lận) để ngưỡng bất thường đủ rộng, giúp mô hình phát hiện được một phần giao dịch gian lận.")

    # ---- 3.8 Tinh chỉnh
    R.h2("Tinh chỉnh siêu tham số")
    lr_th, rf_th = th["logistic_regression"], th["random_forest"]
    R.para("Grid Search kết hợp 3-fold Cross-Validation (CrossValidator của Spark MLlib) được thực hiện với độ "
           "đo areaUnderPR. Lưới tham số: Logistic Regression với regParam ∈ {0,001; 0,01; 0,1} và "
           "elasticNetParam ∈ {0; 0,5}; Random Forest với numTrees ∈ {50; 100} và maxDepth ∈ {5; 10}. Kết quả "
           f"PR-AUC trung bình của từng tổ hợp được thể hiện ở {R.next_figure_ref()}.")
    if tu:
        R.figure("tinh_chinh_luoi_tham_so", "PR-AUC trung bình 3-fold của các tổ hợp tham số")
    strat = tu.get("strategy", {})
    rows = []
    for name, key, params in [("Logistic Regression", "logistic_regression",
                               f"regParam = {lr_th['regParam']}, elasticNetParam = {lr_th['elasticNetParam']}"),
                              ("Random Forest", "random_forest",
                               f"numTrees = {rf_th['numTrees']}, maxDepth = {rf_th['maxDepth']}")]:
        sc_ = strat.get(name, {})
        cmp = (f"class weight: {vn(sc_['class_weight'], 4)} / SMOTE: {vn(sc_['smote'], 4)}" if sc_ else "")
        rows.append([name, params, f"{th[key]['strategy']} ({cmp})" if cmp else th[key]["strategy"],
                     vn(th[key]["threshold"], 4)])
    R.para("Tiếp theo, tập huấn luyện được tách tiếp theo thời gian: 85% đầu để huấn luyện, 15% cuối làm tập "
           "**validation**. Với bộ tham số tốt nhất, mỗi mô hình được huấn luyện theo cả hai chiến lược (class "
           "weight và SMOTE – SMOTE chỉ áp dụng trên phần 85%) và so sánh PR-AUC trên validation để chọn chiến "
           "lược; ngưỡng phân loại được tối ưu (F1 lớn nhất) trên validation bằng chính mô hình của chiến lược "
           "đã chọn, vì xác suất của hai chiến lược không cùng thang đo. Cuối cùng mô hình được huấn luyện lại "
           f"trên toàn bộ tập huấn luyện. Kết quả được tổng hợp trong {R.next_table_ref()}.")
    R.table("Kết quả tinh chỉnh", ["Mô hình", "Tham số tốt nhất", "Chiến lược (PR-AUC validation)", "Ngưỡng"], rows,
            [3.4, 4.6, 5.8, 2.2], ["left", "left", "left", "right"])

    # ---- 3.9 Đánh giá
    R.h2("Đánh giá và so sánh mô hình")
    R.para(f"Các mô hình sau tinh chỉnh được đánh giá trên tập test ({vn(s['test_rows'])} giao dịch, "
           f"{vn(s['test_fraud'])} gian lận) – dữ liệu mà mô hình chưa từng thấy. Kết quả trong "
           f"{R.next_table_ref()}.")
    rows = []
    for name in ["Logistic Regression", "Random Forest", "Isolation Forest"]:
        r = m.loc[name]
        rows.append([name, vn(r["precision"], 4), vn(r["recall"], 4), vn(r["f1_score"], 4),
                     vn(r["pr_auc"], 4), vn(r["roc_auc"], 4)])
    R.table("Kết quả đánh giá 3 mô hình trên tập test",
            ["Mô hình", "Precision", "Recall", "F1-score", "PR-AUC", "ROC-AUC"], rows,
            [4.0, 2.4, 2.4, 2.4, 2.4, 2.4], ["left"] + ["right"] * 5)
    R.figure("so_sanh_do_do", "So sánh các độ đo của 3 mô hình")
    R.para(f"Ma trận nhầm lẫn ({R.next_figure_ref()}) cho thấy rõ sự khác biệt: Random Forest phát hiện đúng "
           f"{int(rf['tp'])}/{int(rf['tp'] + rf['fn'])} giao dịch gian lận với chỉ {int(rf['fp'])} lần báo nhầm "
           f"trên {vn(rf['tn'] + rf['fp'])} giao dịch bình thường; Logistic Regression phát hiện "
           f"{int(lr['tp'])} ca nhưng báo nhầm {vn(lr['fp'])} lần; Isolation Forest báo nhầm tới {vn(iso['fp'])} "
           f"lần mà chỉ phát hiện {int(iso['tp'])} ca.")
    R.figure("ma_tran_nham_lan", "Ma trận nhầm lẫn của 3 mô hình trên tập test")
    R.para(f"Đường Precision–Recall và ROC ({R.next_figure_ref()}) minh họa nhận định ở Chương 2: cả ba mô "
           "hình đều có ROC-AUC cao (trên 0,9), nhưng PR-AUC chênh lệch rất lớn – PR-AUC mới phản ánh đúng khả "
           "năng phát hiện gian lận.")
    R.figure("duong_pr_roc", "Đường Precision–Recall và ROC của 3 mô hình")
    R.para(f"{R.next_figure_ref()} cho thấy tác động của ngưỡng phân loại với Random Forest: tăng ngưỡng làm "
           "Precision tăng và Recall giảm. Ngưỡng đã tinh chỉnh "
           f"({vn(rf_th['threshold'], 3)}) được chọn trên tập validation, không dùng tập test.")
    R.figure("nguong_phan_loai", "Precision, Recall, F1 của Random Forest theo ngưỡng phân loại")
    imp = s["rf_importance_top"]
    top_imp = list(imp)[:5]
    R.para(f"{R.next_figure_ref()} trình bày 15 đặc trưng quan trọng nhất của Random Forest; đứng đầu là "
           + ", ".join(top_imp) + ", khớp với các đặc trưng có tương quan mạnh với nhãn ở phần khảo sát dữ liệu.")
    R.figure("do_quan_trong_dac_trung", "Mức độ quan trọng của đặc trưng trong Random Forest", 14)

    R.h3("Ảnh hưởng của đặc trưng giờ")
    b, a = d["before"], d["after"]
    rows = []
    for name in ["Logistic Regression", "Random Forest", "Isolation Forest"]:
        rows.append([name, f"{vn(b.loc[name, 'f1_score'], 4)} → {vn(a.loc[name, 'f1_score'], 4)}",
                     f"{vn(b.loc[name, 'pr_auc'], 4)} → {vn(a.loc[name, 'pr_auc'], 4)}",
                     f"{vn(b.loc[name, 'recall'], 4)} → {vn(a.loc[name, 'recall'], 4)}"])
    R.para("Để đánh giá khách quan đặc trưng mới, nhóm huấn luyện các mô hình với cùng tham số mặc định và "
           f"ngưỡng 0,5, chỉ thay Time_scaled bằng Hour_sin, Hour_cos ({R.next_table_ref()}, {R.next_figure_ref()}).")
    R.table("So sánh trước và sau khi thay Time_scaled bằng đặc trưng giờ (tham số mặc định, ngưỡng 0,5)",
            ["Mô hình", "F1-score", "PR-AUC", "Recall"], rows, [4.0, 4.0, 4.0, 4.0],
            ["left", "center", "center", "center"])
    R.figure("anh_huong_dac_trung_gio", "Ảnh hưởng của đặc trưng giờ đến các độ đo")
    d_rf = a.loc["Random Forest", "pr_auc"] - b.loc["Random Forest", "pr_auc"]
    R.para("Kết quả cho thấy đặc trưng giờ **gần như không làm thay đổi độ chính xác** của Random Forest "
           f"(PR-AUC thay đổi {vn(d_rf, 4)}): phần lớn thông tin phân biệt gian lận đã nằm trong V1–V28. Tuy "
           "nhiên đặc trưng giờ vẫn được giữ lại vì hai lý do: (1) loại bỏ được đặc trưng Time_scaled không hợp "
           "lệ về mặt phương pháp như đã phân tích; (2) giúp **giải thích kết quả bằng ngôn ngữ dễ hiểu** "
           "(“giao dịch lúc 3 giờ sáng”) trong ứng dụng – điều mà các thành phần ẩn danh V1–V28 không làm được.")

    R.h3("Phân tích và thảo luận")
    R.bullets([
        f"**Random Forest là mô hình tốt nhất** với F1-score {vn(rf['f1_score'], 4)} và PR-AUC "
        f"{vn(rf['pr_auc'], 4)}: cân bằng tốt giữa phát hiện gian lận và hạn chế báo nhầm, nhờ học được quan "
        "hệ phi tuyến giữa các đặc trưng. Mô hình này được chọn cho ứng dụng.",
        f"**Logistic Regression** cho thấy rõ ảnh hưởng của ngưỡng phân loại: ở ngưỡng mặc định 0,5 mô hình "
        f"có Recall {vn(d['after'].loc['Logistic Regression', 'recall'], 4)} nhưng Precision chỉ "
        f"{vn(d['after'].loc['Logistic Regression', 'precision'], 4)} (cảnh báo quá nhiều giao dịch bình "
        f"thường do trọng số lớp rất lớn); sau khi tối ưu ngưỡng, Precision tăng lên {vn(lr['precision'], 4)} "
        f"nhưng Recall giảm còn {vn(lr['recall'], 4)}. Với ranh giới quyết định tuyến tính, mô hình không đạt "
        "được đồng thời cả hai như Random Forest.",
        f"**Isolation Forest** có Precision rất thấp ({vn(iso['precision'], 4)}). Đây không phải lỗi cài đặt "
        "mà là hạn chế đã biết của phương pháp không giám sát: thuật toán tìm các điểm “khác biệt” trong toàn "
        "bộ không gian đặc trưng, trong khi nhiều giao dịch hợp lệ cũng khác biệt (số tiền lớn, hành vi hiếm) "
        "và nhiều giao dịch gian lận lại trông bình thường. Isolation Forest phù hợp khi không có nhãn.",
        "Với dữ liệu mất cân bằng nặng, ROC-AUC của cả ba mô hình đều cao, nhưng PR-AUC phân biệt rõ chất "
        "lượng – khẳng định việc chọn PR-AUC làm độ đo chính là đúng đắn.",
    ])

    # ---- 3.10 Giao dịch nghi vấn
    R.h2("Trực quan hóa giao dịch nghi vấn")
    caught = s["blocked_amount"]
    R.para(f"{R.next_figure_ref()} biểu diễn toàn bộ giao dịch của tập test theo thời điểm và số tiền, phân "
           "loại theo kết quả của Random Forest (ngưỡng đã tinh chỉnh): giao dịch được duyệt đúng, giao dịch "
           "gian lận bị phát hiện, báo nhầm và bỏ sót. Các giao dịch bị phát hiện đúng có tổng giá trị "
           f"**{vn(caught, 2)} €** trên tổng {vn(s['fraud_amount_test'], 2)} € gian lận của tập test "
           f"({pct(caught / s['fraud_amount_test'], 1)} số tiền gian lận được ngăn chặn).")
    R.figure("giao_dich_nghi_van", "Các giao dịch nghi vấn do Random Forest phát hiện trên tập test")

    # ---- 3.11 Ứng dụng
    R.h2("Ứng dụng minh họa")
    R.para("Ứng dụng web minh họa gồm backend **FastAPI** và frontend **Next.js**, được thiết kế cho người dùng "
           "không chuyên (nhân viên ngân hàng, khách hàng):")
    R.bullets([
        "**Tab ETL:** theo dõi quá trình Spark xử lý dữ liệu (các bước, job/stage/task chạy song song) và các "
        "biểu đồ khảo sát dữ liệu, kèm nhật ký sự kiện viết bằng câu dễ hiểu.",
        "**Tab Huấn luyện & Đánh giá:** chạy từng bước xử lý mất cân bằng, huấn luyện, tinh chỉnh, đánh giá và "
        "xem bảng so sánh, ma trận nhầm lẫn.",
        "**Tab Suy luận:** mô phỏng giám sát trực tiếp – phát lại giao dịch của tập test theo thứ tự thời gian; "
        "mô hình tốt nhất chấm điểm rủi ro 0–100 cho từng giao dịch, cảnh báo giao dịch nghi vấn và **giải "
        "thích lý do** bằng SHAP được gộp thành ba nhóm dễ hiểu: hành vi thẻ, giờ giao dịch, số tiền; kèm các "
        "giao dịch tương tự trong quá khứ và chức năng thử thay đổi số tiền/giờ.",
    ])
    R.para("__(Ảnh chụp màn hình giao diện ứng dụng được bổ sung khi hoàn thiện ứng dụng.)__")

    R.h2("Hạn chế của giải pháp")
    R.bullets([
        "Dữ liệu chỉ gồm 2 ngày giao dịch; các đặc trưng V1–V28 đã ẩn danh nên không thể tạo thêm đặc trưng "
        "nghiệp vụ (lịch sử chi tiêu theo thẻ, vị trí, thiết bị…).",
        "Giờ trong ngày được suy ra từ giả định giao dịch đầu tiên lúc 0h.",
        "HDFS và Spark chạy trên một máy (pseudo-cluster), chưa đánh giá hiệu năng trên cụm nhiều máy.",
        "Xử lý theo lô; chưa triển khai xử lý luồng thời gian thực bằng Spark Structured Streaming.",
    ])


def conclusion(R: Report, d: dict):
    m = d["m"]
    rf = m.loc["Random Forest"]
    R.unnumbered("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    R.para("**Kết quả đạt được.** Đề tài đã xây dựng hoàn chỉnh quy trình phát hiện giao dịch tài chính bất "
           "thường trên nền tảng Big Data:")
    R.bullets([
        "Lưu trữ dữ liệu trên HDFS (pseudo-cluster bằng Docker) và xử lý bằng Spark DataFrame song song: khảo "
        "sát, làm sạch (loại trùng lặp, phân tích ngoại lệ), tạo đặc trưng giờ, chuẩn hóa và chia tập theo "
        "thời gian.",
        "Xử lý dữ liệu mất cân bằng bằng class weight và SMOTE; huấn luyện, tinh chỉnh (Grid Search + "
        "Cross-Validation, tối ưu ngưỡng) và so sánh 3 mô hình.",
        f"Mô hình tốt nhất – Random Forest – đạt Precision {vn(rf['precision'], 4)}, Recall "
        f"{vn(rf['recall'], 4)}, F1-score {vn(rf['f1_score'], 4)}, PR-AUC {vn(rf['pr_auc'], 4)} trên tập test.",
        "Trực quan hóa dữ liệu rút ra các thông tin có giá trị: khung giờ ban đêm và một số khoảng số tiền có "
        "rủi ro gian lận cao vượt trội; xây dựng ứng dụng minh họa có giải thích kết quả.",
    ])
    R.para("**Hướng phát triển.**")
    R.bullets([
        "Triển khai xử lý luồng thời gian thực với Kafka và Spark Structured Streaming để chấm điểm giao dịch "
        "ngay khi phát sinh.",
        "Triển khai trên cụm Hadoop/Spark nhiều máy và đánh giá khả năng mở rộng với dữ liệu lớn hơn.",
        "Thử nghiệm các mô hình mạnh hơn như Gradient-Boosted Trees, XGBoost, LightGBM và các mô hình học sâu.",
        "Bổ sung đặc trưng hành vi theo từng thẻ (tần suất, số tiền trung bình gần đây) khi có dữ liệu chưa ẩn "
        "danh; cập nhật mô hình định kỳ để thích ứng với thủ đoạn gian lận mới.",
    ])


def references(R: Report):
    R.unnumbered("TÀI LIỆU THAM KHẢO")
    refs = [
        "Machine Learning Group – ULB, “Credit Card Fraud Detection,” Kaggle. [Online]. Available: "
        "https://www.kaggle.com/datasets/mlg-ulb/creditcardfraud",
        "A. Dal Pozzolo, O. Caelen, R. A. Johnson, G. Bontempi, “Calibrating Probability with Undersampling "
        "for Unbalanced Classification,” in Proc. IEEE Symposium Series on Computational Intelligence (SSCI), 2015.",
        "K. Shvachko, H. Kuang, S. Radia, R. Chansler, “The Hadoop Distributed File System,” in Proc. IEEE 26th "
        "Symposium on Mass Storage Systems and Technologies (MSST), 2010.",
        "M. Zaharia et al., “Resilient Distributed Datasets: A Fault-Tolerant Abstraction for In-Memory Cluster "
        "Computing,” in Proc. USENIX NSDI, 2012.",
        "M. Zaharia et al., “Apache Spark: A Unified Engine for Big Data Processing,” Communications of the ACM, "
        "vol. 59, no. 11, pp. 56–65, 2016.",
        "N. V. Chawla, K. W. Bowyer, L. O. Hall, W. P. Kegelmeyer, “SMOTE: Synthetic Minority Over-sampling "
        "Technique,” Journal of Artificial Intelligence Research, vol. 16, pp. 321–357, 2002.",
        "L. Breiman, “Random Forests,” Machine Learning, vol. 45, no. 1, pp. 5–32, 2001.",
        "F. T. Liu, K. M. Ting, Z.-H. Zhou, “Isolation Forest,” in Proc. IEEE International Conference on Data "
        "Mining (ICDM), 2008.",
        "J. Davis, M. Goadrich, “The Relationship Between Precision-Recall and ROC Curves,” in Proc. ICML, 2006.",
        "T. Saito, M. Rehmsmeier, “The Precision-Recall Plot Is More Informative than the ROC Plot When "
        "Evaluating Binary Classifiers on Imbalanced Datasets,” PLOS ONE, vol. 10, no. 3, 2015.",
        "S. M. Lundberg, S.-I. Lee, “A Unified Approach to Interpreting Model Predictions,” in Advances in "
        "Neural Information Processing Systems (NeurIPS), 2017.",
        "S. M. Lundberg et al., “From Local Explanations to Global Understanding with Explainable AI for "
        "Trees,” Nature Machine Intelligence, vol. 2, pp. 56–67, 2020.",
        "Apache Software Foundation, “Spark 4 Documentation – MLlib Guide.” [Online]. Available: "
        "https://spark.apache.org/docs/latest/ml-guide.html",
        "Apache Software Foundation, “HDFS Architecture.” [Online]. Available: https://hadoop.apache.org/docs/"
        "stable/hadoop-project-dist/hadoop-hdfs/HdfsDesign.html",
    ]
    for i, ref in enumerate(refs, start=1):
        p = R.doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.left_indent = Cm(0.9)
        p.add_run(f"[{i}]\t{ref}")


def update_fields_with_word(path: Path) -> None:
    script = REPORTS_DIR / "update_fields.ps1"
    if sys.platform != "win32" or not script.exists():
        print("[Report] Mở file bằng Word, nhấn Ctrl+A rồi F9 để cập nhật mục lục và danh mục.")
        return
    try:
        out = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script),
                              "-Path", str(path)], capture_output=True, text=True, timeout=300)
        print(out.stdout.strip() or out.stderr.strip())
    except Exception as e:  # Word không cài / lỗi COM - không chặn việc sinh báo cáo
        print(f"[Report] Không cập nhật được mục lục tự động ({e}). Mở Word, Ctrl+A rồi F9.")


def main():
    d = load_inputs()
    R = Report(TEMPLATE_PATH)
    chapter1(R, d)
    chapter2(R, d)
    chapter3(R, d)
    conclusion(R, d)
    references(R)
    R.save(OUTPUT_PATH)
    print(f"[Report] Đã sinh báo cáo -> {OUTPUT_PATH}")
    update_fields_with_word(OUTPUT_PATH)


if __name__ == "__main__":
    main()
