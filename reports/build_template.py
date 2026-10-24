"""
Dựng khung báo cáo (reports/template/report_template.docx) từ file Word mẫu định dạng của khoa.

Chỉ lấy ĐỊNH DẠNG của file mẫu (trang bìa có logo + viền trang, kiểu chữ, kiểu chương/mục h1-h4,
chú thích bảng/hình, mục lục, đánh số trang La Mã/Ả Rập theo section); toàn bộ nội dung cũ bị
xóa. Trang bìa và bảng phân công được điền thông tin của đồ án này.

Chạy 1 lần (hoặc khi đổi thông tin trang bìa):
    python reports/build_template.py "<file mẫu>.docx"
Sau đó: python reports/generate_report.py
"""

import copy
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_PATH = ROOT_DIR / "reports" / "template" / "report_template.docx"

# Thông tin trang bìa / bảng phân công. Tên, MSSV sinh viên đang để TẠM - sửa lại cho đúng.
COURSE_TITLE = "ĐỒ ÁN MÔN HỌC NHẬP MÔN BIG DATA"
TOPIC_TITLE = "PHÁT HIỆN GIAO DỊCH TÀI CHÍNH BẤT THƯỜNG"
LECTURER = "TS. Ngô Dương Hà"
GROUP_LABEL = "Nhóm …"
PLACE_DATE = "TP. Hồ Chí Minh, tháng 10 - 2026"
MEMBERS = [
    {
        "mssv": "2001xxxxxx",
        "name": "Nguyễn Văn A",
        "leader": True,
        "tasks": "Quản lý tiến độ đồ án. Triển khai HDFS (Docker) và Spark; xây dựng quy trình "
                 "ETL, làm sạch, tiền xử lý và tạo đặc trưng; viết Chương 1, Chương 2.",
    },
    {
        "mssv": "2001xxxxxx",
        "name": "Trần Thị B",
        "leader": False,
        "tasks": "Xử lý dữ liệu mất cân bằng; huấn luyện, tinh chỉnh và đánh giá mô hình; trực "
                 "quan hóa dữ liệu, xây dựng ứng dụng minh họa; viết Chương 3, slide thuyết trình.",
    },
]

TOC_INSTR = ' TOC \\h \\z \\t "h1,1,h2,2,h3,3,h4,4,More,1" '
MARKER_TOF_FIGURES = "{{DANH_MUC_HINH}}"
MARKER_TOF_TABLES = "{{DANH_MUC_BANG}}"


def set_text(paragraph: Paragraph, text: str) -> None:
    """Đặt lại nội dung đoạn văn, giữ định dạng của run đầu tiên."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def find_paragraph(body_children, doc, startswith: str) -> int:
    for i, el in enumerate(body_children):
        if el.tag == qn("w:p") and Paragraph(el, doc).text.strip().startswith(startswith):
            return i
    raise ValueError(f"Không tìm thấy đoạn bắt đầu bằng '{startswith}' trong file mẫu")


def field_paragraph(instr: str, placeholder: str, style_el=None):
    """Tạo 1 đoạn chứa trường (field) Word, đánh dấu dirty để Word tự cập nhật khi mở."""
    from docx.oxml import parse_xml
    xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:pPr><w:ind w:firstLine="0"/></w:pPr>'
        '<w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
        f'<w:r><w:instrText xml:space="preserve">{instr}</w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r><w:t>{placeholder}</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
    )
    return parse_xml(xml)


def marker_paragraph(text: str):
    from docx.oxml import parse_xml
    return parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:pPr><w:ind w:firstLine="0"/></w:pPr><w:r><w:t>{text}</w:t></w:r></w:p>'
    )


def fill_cover(doc, children) -> None:
    i_course = find_paragraph(children, doc, "ĐỒ ÁN")
    set_text(Paragraph(children[i_course], doc), COURSE_TITLE)

    i_topic_label = find_paragraph(children, doc, "ĐỀ TÀI")
    topic = Paragraph(children[i_topic_label + 1], doc)
    set_text(topic, TOPIC_TITLE)

    i_students = find_paragraph(children, doc, "Sinh viên thực hiện")
    students_label = Paragraph(children[i_students], doc)
    # Dòng giảng viên: nhân bản dòng "Sinh viên thực hiện" để giữ đúng định dạng.
    lecturer_el = copy.deepcopy(children[i_students])
    children[i_students].addprevious(lecturer_el)
    lecturer = Paragraph(lecturer_el, doc)
    lecturer.runs[0].text = "Giảng viên hướng dẫn: "
    for r in lecturer.runs[1:]:
        r._r.getparent().remove(r._r)
    bold = lecturer.add_run(LECTURER)
    bold.bold = True

    students_label.runs[0].text = "Sinh viên thực hiện: "
    for r in students_label.runs[1:]:
        r._r.getparent().remove(r._r)
    students_label.add_run(GROUP_LABEL).bold = True

    # Danh sách sinh viên: các đoạn "List Paragraph" ngay sau dòng nhãn.
    el = students_label._p.getnext()
    member_els = []
    while el is not None and el.tag == qn("w:p") and Paragraph(el, doc).style.name == "List Paragraph":
        member_els.append(el)
        el = el.getnext()
    for el, m in zip(member_els, MEMBERS):
        suffix = " (NT)" if m["leader"] else ""
        set_text(Paragraph(el, doc), f"{m['mssv']} - {m['name']}{suffix}")
    for el in member_els[len(MEMBERS):]:
        el.getparent().remove(el)

    i_date = find_paragraph(list(doc.element.body.iterchildren()), doc, "TP. Hồ Chí Minh")
    date = Paragraph(list(doc.element.body.iterchildren())[i_date], doc)
    set_text(date, PLACE_DATE)
    # Ngắt trang sau dòng ngày tháng nằm trong run bị xóa ở set_text -> thêm lại.
    date.add_run().add_break(WD_BREAK.PAGE)


def fill_member_table(doc, children) -> None:
    i_title = find_paragraph(children, doc, "BẢNG PHÂN CÔNG")
    group = Paragraph(children[i_title + 1], doc)
    set_text(group, GROUP_LABEL.upper())
    table = Table(children[i_title + 2], doc)
    data_rows = table.rows[1:]
    for idx, (row, m) in enumerate(zip(data_rows, MEMBERS), start=1):
        values = [str(idx), m["mssv"], m["name"], m["tasks"], ""]
        for cell, value in zip(row.cells, values):
            for p in cell.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            set_text(cell.paragraphs[0], value)
    for row in data_rows[len(MEMBERS):]:
        row._tr.getparent().remove(row._tr)


def replace_field_block(doc, children, heading_text: str, new_el) -> None:
    """Thay các đoạn mục lục/danh mục cũ (ngay sau tiêu đề `heading_text`, tới ngắt trang hoặc
    ngắt section tiếp theo) bằng `new_el`."""
    i = find_paragraph(children, doc, heading_text)
    el = children[i].getnext()
    to_remove = []
    while el is not None:
        is_break = el.xpath('.//w:br[@w:type="page"]') or el.xpath("./w:pPr/w:sectPr")
        if is_break:
            break
        to_remove.append(el)
        el = el.getnext()
    children[i].addnext(new_el)
    for e in to_remove:
        e.getparent().remove(e)


def strip_body(doc) -> None:
    """Xóa toàn bộ nội dung sau section phần đầu (mục lục, danh mục), giữ sectPr cuối."""
    body = doc.element.body
    children = list(body.iterchildren())
    i_tables_heading = find_paragraph(children, doc, "DANH MỤC BẢNG BIỂU")
    # Đoạn chứa sectPr đầu tiên sau "DANH MỤC BẢNG BIỂU" là ranh giới section phần đầu / nội dung.
    boundary = next(
        i for i in range(i_tables_heading, len(children))
        if children[i].tag == qn("w:p") and children[i].xpath("./w:pPr/w:sectPr")
    )
    for el in children[boundary + 1:]:
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def drop_unused_images(doc) -> None:
    part = doc.part
    used = set(doc.element.xpath("//@r:embed"))
    for rId, rel in list(part.rels.items()):
        if rel.reltype.endswith("/image") and rId not in used:
            part.drop_rel(rId)


def main(template_path: str) -> None:
    doc = Document(template_path)
    body = doc.element.body

    fill_cover(doc, list(body.iterchildren()))
    fill_member_table(doc, list(body.iterchildren()))

    replace_field_block(doc, list(body.iterchildren()), "MỤC LỤC",
                        field_paragraph(TOC_INSTR, "Nhấn chuột phải → Update Field để cập nhật mục lục"))
    replace_field_block(doc, list(body.iterchildren()), "DANH MỤC HÌNH ẢNH",
                        marker_paragraph(MARKER_TOF_FIGURES))
    replace_field_block(doc, list(body.iterchildren()), "DANH MỤC BẢNG BIỂU",
                        marker_paragraph(MARKER_TOF_TABLES))
    strip_body(doc)
    drop_unused_images(doc)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"[Template] Đã tạo khung báo cáo -> {OUTPUT_PATH} "
          f"({OUTPUT_PATH.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Cách dùng: python reports/build_template.py <file Word mẫu>.docx")
        sys.exit(1)
    main(sys.argv[1])
